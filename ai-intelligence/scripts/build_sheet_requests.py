#!/usr/bin/env python3
"""Emit Google Sheets updateCells requests from the staged workbook."""
from __future__ import annotations

import argparse
import json
from openpyxl import load_workbook


SHEETS = {
    "Overview": 2037736221,
    "Master Directory": 1442268904,
    "APAC and SEA": 520206069,
    "Map Shortlist": 700286018,
    "Layer Taxonomy": 234483116,
    "Six-Layer Stack": 1684202609,
    "Landscape Panels": 1774202609,
    "Application Landscape": 1472528957,
    "Lifecycle Archive": 1872528957,
    "Companies DB": 192517072,
    "Leadership DB": 437031203,
    "Market Data": 609029537,
    "Sources DB": 877681264,
    "Taxonomy v2": 584434877,
    "Coverage QA": 1329961750,
    "Change Log": 829901643,
    "Funding Rounds": 837298840,
    "Logo Registry": 2100000001,
    "Data Quality": 2100000002,
    "Review Queue": 2100000003,
}


def user_value(value):
    if value is None:
        return {"stringValue": ""}
    if isinstance(value, bool):
        return {"boolValue": value}
    if isinstance(value, (int, float)):
        return {"numberValue": value}
    text = str(value)
    if text.startswith("="):
        return {"formulaValue": text}
    return {"stringValue": text}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--sheet", required=True, choices=SHEETS)
    ap.add_argument("--start-row", type=int, default=2)
    ap.add_argument("--end-row", type=int)
    ap.add_argument("--columns", help="Comma-separated header names to write as separate one-column requests")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, data_only=False)
    ws = wb[args.sheet]
    start = max(1, args.start_row)
    end = min(args.end_row or ws.max_row, ws.max_row)
    requests = []
    if args.columns:
        header_map = {c.value: c.column for c in ws[1]}
        for key in [v.strip() for v in args.columns.split(",") if v.strip()]:
            col = header_map[key]
            rows = [{"values": [{"userEnteredValue": user_value(ws.cell(r, col).value)}]} for r in range(start, end + 1)]
            requests.append({"updateCells": {"start": {"sheetId": SHEETS[args.sheet], "rowIndex": start - 1, "columnIndex": col - 1}, "rows": rows, "fields": "userEnteredValue"}})
    else:
        rows = []
        for row in ws.iter_rows(min_row=start, max_row=end, max_col=ws.max_column, values_only=True):
            rows.append({"values": [{"userEnteredValue": user_value(v)} for v in row]})
        requests.append({"updateCells": {"start": {"sheetId": SHEETS[args.sheet], "rowIndex": start - 1, "columnIndex": 0}, "rows": rows, "fields": "userEnteredValue"}})
    print(json.dumps({"requests": requests}, ensure_ascii=False))


if __name__ == "__main__":
    main()
