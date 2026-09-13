#!/usr/bin/env python3
"""Apply source-backed company and leadership X accounts to the workbook."""

from __future__ import annotations

import argparse
import copy
import json
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

import openpyxl


def clean(value) -> str:
    return str(value or "").strip()


def canonical_x(value) -> str:
    try:
        parsed = urlparse(clean(value))
        if parsed.netloc.lower().removeprefix("www.") not in {"x.com", "twitter.com", "mobile.twitter.com"}:
            return ""
        handle = parsed.path.strip("/").split("/")[0]
        if not handle or not handle.replace("_", "").isalnum() or len(handle) > 15:
            return ""
        return f"https://x.com/{handle}"
    except Exception:
        return ""


def headers(sheet) -> dict[str, int]:
    return {clean(cell.value): cell.column for cell in sheet[1] if clean(cell.value)}


def ensure_columns(sheet, names: list[str]) -> dict[str, int]:
    index = headers(sheet)
    template_col = sheet.max_column
    for name in names:
        if name in index:
            continue
        column = sheet.max_column + 1
        source = sheet.cell(1, template_col)
        target = sheet.cell(1, column, name)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)
        sheet.column_dimensions[target.column_letter].width = 24
        index[name] = column
    return index


def rows_by_id(sheet, id_column: int) -> dict[str, int]:
    return {clean(sheet.cell(row, id_column).value): row for row in range(2, sheet.max_row + 1) if clean(sheet.cell(row, id_column).value)}


def next_source_id(sheet, index: dict[str, int]) -> int:
    maximum = 0
    for row in range(2, sheet.max_row + 1):
        value = clean(sheet.cell(row, index["source_id"]).value)
        digits = "".join(character for character in value if character.isdigit())
        if digits:
            maximum = max(maximum, int(digits))
    return maximum + 1


def next_change_id(sheet, index: dict[str, int]) -> int:
    maximum = 0
    for row in range(2, sheet.max_row + 1):
        value = clean(sheet.cell(row, index["change_id"]).value)
        digits = "".join(character for character in value if character.isdigit())
        if digits:
            maximum = max(maximum, int(digits))
    return maximum + 1


def set_if_changed(sheet, row: int, column: int, value) -> tuple[str, str] | None:
    previous = clean(sheet.cell(row, column).value)
    current = clean(value)
    if previous == current:
        return None
    sheet.cell(row, column, value)
    return previous, current


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--research", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    research = json.loads(Path(args.research).read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(args.input)
    companies = workbook["Companies DB"]
    leaders = workbook["Leadership DB"]
    sources = workbook["Sources DB"]
    changes = workbook["Change Log"]
    coverage = workbook["Coverage QA"]

    company_index = ensure_columns(companies, [
        "company_x_url", "company_x_handle", "company_x_source_url",
        "company_x_avatar_url", "company_x_last_verified_at", "company_x_verification_status",
    ])
    leader_index = ensure_columns(leaders, [
        "x_url", "x_handle", "x_source_url", "x_avatar_url", "x_last_verified_at", "x_verification_status",
        "photo_verification_status",
    ])
    source_index = headers(sources)
    change_index = headers(changes)
    company_rows = rows_by_id(companies, company_index["company_id"])
    leader_rows = rows_by_id(leaders, leader_index["company_id"])
    today = date.today().isoformat()
    source_number = next_source_id(sources, source_index)
    change_number = next_change_id(changes, change_index)

    existing_source_urls = {
        clean(sources.cell(row, source_index["source_url"]).value)
        for row in range(2, sources.max_row + 1)
    }
    existing_change_keys = {
        (
            clean(changes.cell(row, change_index["company_id"]).value),
            clean(changes.cell(row, change_index["field_changed"]).value),
            clean(changes.cell(row, change_index["new_value"]).value),
        )
        for row in range(2, changes.max_row + 1)
    }

    counts = {"company": 0, "leader": 0, "misfiled_linkedin_cleared": 0, "sources_added": 0, "changes_added": 0}

    def log_change(company_id: str, field: str, previous: str, current: str, reason: str, source_url: str) -> None:
        nonlocal change_number
        key = (company_id, field, current)
        if key in existing_change_keys:
            return
        row = changes.max_row + 1
        values = {
            "change_id": f"CHG-{change_number:05d}", "changed_at": today, "company_id": company_id,
            "field_changed": field, "previous_value": previous or None, "new_value": current or None,
            "reason": reason, "source_url": source_url, "review_status": "Applied — source-backed",
        }
        for key_name, value in values.items():
            changes.cell(row, change_index[key_name], value)
        change_number += 1
        counts["changes_added"] += 1
        existing_change_keys.add(key)

    def add_source(company_id: str, company_name: str, scope: str, account, source_url: str) -> None:
        nonlocal source_number
        evidence_url = clean(source_url) or clean(account.get("url"))
        unique_url = clean(account.get("url"))
        if unique_url in existing_source_urls:
            return
        row = sources.max_row + 1
        values = {
            "source_id": f"SRC-X-{source_number:05d}", "company_id": company_id,
            "company_name": company_name, "source_type": "Official X account",
            "publisher": "X", "source_title": f"{scope.title()} X profile: @{account.get('handle')}",
            "source_url": unique_url, "publication_date": None, "accessed_at": today,
            "reliability_tier": "Tier 1 — first-party identity link" if "official" in clean(account.get("sourceType")) or clean(account.get("sourceType")) == "reviewed_leadership_identity" else "Tier 2 — structured identity record",
            "fields_supported": f"{scope}_x_url | {scope}_x_handle",
            "notes": f"Validated with XActions public profile reader; identity evidence: {', '.join(account.get('reasons') or [])}. First-party evidence: {evidence_url}",
        }
        for key_name, value in values.items():
            sources.cell(row, source_index[key_name], value)
        source_number += 1
        counts["sources_added"] += 1
        existing_source_urls.add(unique_url)

    for record in research.get("records", []):
        company_id = clean(record.get("companyId"))
        company_name = clean(record.get("companyName"))
        accepted = record.get("accepted") or {}

        if company_id in company_rows and accepted.get("company"):
            account = accepted["company"]
            row = company_rows[company_id]
            url = canonical_x(account.get("url"))
            values = {
                "company_x_url": url, "company_x_handle": f"@{account.get('handle')}",
                "company_x_source_url": account.get("sourceUrl"), "company_x_last_verified_at": today,
                "company_x_avatar_url": clean((account.get("profile") or {}).get("avatar")),
                "company_x_verification_status": f"Verified {account.get('confidence')} confidence — X profile active and identity matched",
            }
            for field, value in values.items():
                changed = set_if_changed(companies, row, company_index[field], value)
                if changed:
                    log_change(company_id, field, *changed, "Official X account verified from first-party link and public profile identity", account.get("sourceUrl"))
            add_source(company_id, company_name, "company", account, account.get("sourceUrl"))
            counts["company"] += 1

        if company_id in leader_rows:
            row = leader_rows[company_id]
            linkedin_value = clean(leaders.cell(row, leader_index["linkedin_url"]).value)
            if canonical_x(linkedin_value):
                leaders.cell(row, leader_index["linkedin_url"], "Not publicly disclosed")
                counts["misfiled_linkedin_cleared"] += 1
                log_change(company_id, "leadership.linkedin_url", linkedin_value, "Not publicly disclosed", "X URL removed from LinkedIn-only field", linkedin_value)

            if accepted.get("leader"):
                account = accepted["leader"]
                url = canonical_x(account.get("url"))
                values = {
                    "x_url": url, "x_handle": f"@{account.get('handle')}",
                    "x_source_url": account.get("sourceUrl"), "x_last_verified_at": today,
                    "x_avatar_url": clean((account.get("profile") or {}).get("avatar")),
                    "x_verification_status": f"Verified {account.get('confidence')} confidence — named person and organisation context matched",
                }
                avatar = clean((account.get("profile") or {}).get("avatar"))
                if avatar and not clean(leaders.cell(row, leader_index["photo_url"]).value).startswith("http"):
                    values.update({
                        "photo_url": avatar,
                        "photo_source_url": url,
                        "photo_verification_status": "Verified X profile portrait — named person and current organisation matched",
                    })
                for field, value in values.items():
                    changed = set_if_changed(leaders, row, leader_index[field], value)
                    if changed:
                        log_change(company_id, f"leadership.{field}", *changed, "Leadership X account verified against person identity and current organisation context", account.get("sourceUrl"))
                add_source(company_id, company_name, "leader", account, account.get("sourceUrl"))
                counts["leader"] += 1

    # Append/update evidence-coverage metrics without claiming unresolved accounts are complete.
    coverage_index = headers(coverage)
    metrics = {
        "Verified official company X accounts": (counts["company"], companies.max_row - 1, "First-party page plus active X profile identity match"),
        "Verified leadership X accounts": (counts["leader"], leaders.max_row - 1, "Named-person identity plus first-party/Wikidata/company-context evidence"),
    }
    metric_rows = rows_by_id(coverage, coverage_index["metric"])
    for metric, (current, target, note) in metrics.items():
        row = metric_rows.get(metric) or coverage.max_row + 1
        values = {
            "metric": metric, "current_count": current, "target_count": target,
            "coverage_rate": current / target if target else 0,
            "status": "Evidence coverage", "notes": note,
        }
        for key_name, value in values.items():
            coverage.cell(row, coverage_index[key_name], value)
        coverage.cell(row, coverage_index["coverage_rate"]).number_format = "0.0%"

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({"output": str(output), **counts}, indent=2))


if __name__ == "__main__":
    main()
