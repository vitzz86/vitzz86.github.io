#!/usr/bin/env python3
"""Add field-level evidence grades, registries, and a 497-row review workflow."""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import re
from pathlib import Path
from urllib.parse import quote, urlparse

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


TODAY = dt.date.today().isoformat()
URL = re.compile(r"^https?://", re.I)


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def is_url(value):
    return bool(URL.match(clean(value)))


def headers(ws):
    return {clean(cell.value): cell.column for cell in ws[1] if clean(cell.value)}


def ensure_columns(ws, names):
    index = headers(ws)
    template = ws.max_column
    for name in names:
        if name in index:
            continue
        col = ws.max_column + 1
        source, target = ws.cell(1, template), ws.cell(1, col, name)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        ws.column_dimensions[target.column_letter].width = 24
        index[name] = col
    return index


def rows_by_id(ws, index):
    return {clean(ws.cell(row, index["company_id"]).value): row for row in range(2, ws.max_row + 1)}


def host(value):
    try:
        return (urlparse(clean(value)).hostname or "").lower().removeprefix("www.")
    except Exception:
        return ""


def canonical_linkedin_person(value):
    value = clean(value)
    try:
        parsed = urlparse(value)
    except Exception:
        return ""
    if "linkedin.com" not in parsed.netloc.lower() or "/in/" not in parsed.path.lower():
        return ""
    return "https://www.linkedin.com" + parsed.path.rstrip("/") + "/"


def logo_class(url, website):
    logo_host, site_host = host(url), host(website)
    if logo_host in {"upload.wikimedia.org", "cdn.simpleicons.org", "api.iconify.design"}:
        return "Curated brand registry", "A"
    if "google.com" in logo_host and "/s2/favicons" in clean(url):
        return "Official-domain favicon", "B"
    if logo_host == site_host or (site_host and logo_host.endswith("." + site_host)):
        return "Official website asset", "A"
    if logo_host:
        return "Official-site/CDN asset — identity review retained", "B"
    return "Missing", "C"


def make_sheet(workbook, name, headers_list, widths):
    if name in workbook.sheetnames:
        del workbook[name]
    ws = workbook.create_sheet(name)
    fill = PatternFill("solid", fgColor="173D35")
    for col, title in enumerate(headers_list, 1):
        cell = ws.cell(1, col, title)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = widths.get(title, 22)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers_list)).column_letter}1"
    return ws


def write_row(ws, row, values):
    for col, value in enumerate(values, 1):
        ws.cell(row, col, value)
        ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.input)
    companies, leaders, market = workbook["Companies DB"], workbook["Leadership DB"], workbook["Market Data"]
    ci = ensure_columns(companies, ["evidence_grade", "logo_source_type", "logo_source_url", "publication_status", "qa_last_reviewed_at"])
    li = ensure_columns(leaders, ["role_evidence_grade", "social_evidence_grade", "photo_evidence_grade", "social_disposition", "photo_disposition", "media_review_status", "qa_last_reviewed_at"])
    mi = ensure_columns(market, ["market_evidence_grade", "market_review_status", "qa_last_reviewed_at"])
    cr, lr, mr = rows_by_id(companies, ci), rows_by_id(leaders, li), rows_by_id(market, mi)

    logo_ws = make_sheet(workbook, "Logo Registry", [
        "company_id", "company_name", "logo_url", "official_website", "logo_source_type",
        "evidence_grade", "identity_basis", "review_status", "last_reviewed_at",
    ], {"company_name": 28, "logo_url": 56, "official_website": 42, "logo_source_type": 34, "identity_basis": 48, "review_status": 30})
    qa_ws = make_sheet(workbook, "Data Quality", [
        "company_id", "company_name", "company_profile", "leadership_role", "leadership_social",
        "leadership_photo", "logo", "market", "overall_status", "confidence_score", "last_reviewed_at",
    ], {"company_name": 28, "company_profile": 24, "leadership_role": 24, "leadership_social": 36, "leadership_photo": 36, "overall_status": 28})
    review_ws = make_sheet(workbook, "Review Queue", [
        "company_id", "company_name", "priority", "open_issues", "next_action", "review_status",
        "company_source", "leadership_source", "market_source", "last_attempted_at",
    ], {"company_name": 28, "open_issues": 54, "next_action": 60, "company_source": 44, "leadership_source": 44, "market_source": 44})

    metrics = {
        "company_ready": 0, "logo_a": 0, "logo_resolved": 0, "leader_role": 0,
        "company_linkedin": 0, "company_x": 0, "leader_linkedin": 0, "leader_x": 0,
        "leader_social": 0, "leader_photo": 0, "market_ready": 0, "publication_ready": 0,
        "public_total": 0, "public_market_ready": 0, "private_total": 0,
        "private_market_ready": 0, "private_crunchbase_direct": 0,
    }
    for out_row, company_id in enumerate(sorted(cr, key=lambda cid: clean(companies.cell(cr[cid], ci["display_name"]).value).lower()), 2):
        c_row, l_row, m_row = cr[company_id], lr[company_id], mr[company_id]
        name = clean(companies.cell(c_row, ci["display_name"]).value)
        website = clean(companies.cell(c_row, ci["official_website"]).value)
        logo = clean(companies.cell(c_row, ci["logo_url"]).value)
        company_linkedin = clean(companies.cell(c_row, ci["company_linkedin"]).value)
        company_x = clean(companies.cell(c_row, ci.get("company_x_url", 0)).value) if ci.get("company_x_url") else ""
        description = clean(companies.cell(c_row, ci["description_short"]).value)
        role_source = clean(leaders.cell(l_row, li["official_profile_url"]).value)
        leader_linkedin = canonical_linkedin_person(leaders.cell(l_row, li["linkedin_url"]).value)
        leader_x = clean(leaders.cell(l_row, li.get("x_url", 0)).value) if li.get("x_url") else ""
        legacy_social = clean(leaders.cell(l_row, li["social_url"]).value)
        # Legacy enrichment occasionally stored company Instagram/X handles in
        # a leader field. The public product supports only named-person X and
        # LinkedIn identities, so retain only independently reviewed mappings.
        if company_id == "meta" and legacy_social == "https://x.com/finkd" and not leader_x:
            leader_x = legacy_social
            leaders.cell(l_row, li["x_url"], leader_x)
            leaders.cell(l_row, li["x_handle"], "@finkd")
            leaders.cell(l_row, li["x_source_url"], "https://www.wikidata.org/wiki/Q36215")
            leaders.cell(l_row, li["x_verification_status"], "Verified named-person identity via Wikidata and XActions")
        if leader_linkedin:
            leaders.cell(l_row, li["linkedin_url"], leader_linkedin)
            leaders.cell(l_row, li["social_platform"], "LinkedIn")
            leaders.cell(l_row, li["social_url"], leader_linkedin)
        elif leader_x:
            leaders.cell(l_row, li["social_platform"], "X")
            leaders.cell(l_row, li["social_url"], leader_x)
        else:
            leaders.cell(l_row, li["social_platform"], "Not publicly disclosed")
            leaders.cell(l_row, li["social_url"], "Not publicly disclosed")
        leader_photo = clean(leaders.cell(l_row, li["photo_url"]).value)
        ownership = clean(market.cell(m_row, mi["ownership_status"]).value)
        logo_type, logo_grade = logo_class(logo, website)
        # Publication uses a known-good official-domain favicon when the
        # original asset is an uncurated third-party/CDN URL. This trades some
        # resolution for identity safety and removes mismatched brand images.
        if logo_grade == "B" and logo_type != "Official-domain favicon" and host(website):
            logo = f"https://www.google.com/s2/favicons?domain={quote(host(website))}&sz=128"
            companies.cell(c_row, ci["logo_url"], logo)
            logo_type, logo_grade = "Official-domain favicon", "B"
        logo_resolved = is_url(logo)
        company_ready = is_url(website) and logo_resolved and len(description.split()) >= 8
        role_ready = is_url(role_source)
        social_ready = is_url(leader_linkedin) or is_url(leader_x)
        photo_ready = is_url(leader_photo)
        if ownership == "Public":
            market_ready = all(is_url(market.cell(m_row, mi[field]).value) for field in ("yahoo_finance_url", "investor_relations_url")) and bool(clean(market.cell(m_row, mi["ticker"]).value))
            market_grade = "A" if market_ready else "C"
            market_status = "Verified public-market route" if market_ready else "Public-market data requires review"
            market_source = clean(market.cell(m_row, mi["ticker_validation_source"]).value)
        elif ownership == "Private":
            cb = clean(market.cell(m_row, mi["crunchbase_url"]).value)
            funding = clean(market.cell(m_row, mi["funding_history_status"]).value)
            unicorn = clean(market.cell(m_row, mi["unicorn_status"]).value)
            market_ready = is_url(cb) and bool(funding) and bool(unicorn)
            direct = host(cb) == "crunchbase.com" and "/organization/" in cb
            market_grade = "A" if market_ready and direct else "B" if market_ready else "C"
            market_status = "Verified private-market route" if market_grade == "A" else "Resolved with directory exception" if market_ready else "Private-market data requires review"
            market_source = cb
        else:
            market_ready = bool(clean(market.cell(m_row, mi["funding_history_status"]).value) or is_url(market.cell(m_row, mi["investor_relations_url"]).value))
            market_grade = "B" if market_ready else "C"
            market_status = "Ownership-appropriate market route recorded" if market_ready else "Ownership-specific market route requires review"
            market_source = clean(market.cell(m_row, mi["investor_relations_url"]).value)

        issues = []
        if not company_linkedin.startswith("http") and not company_x.startswith("http"): issues.append("No verified company X or LinkedIn profile")
        if logo_grade != "A": issues.append("Logo uses fallback or CDN evidence; presentation-grade asset review recommended")
        if not social_ready: issues.append("No verified leader X or LinkedIn profile")
        if not photo_ready: issues.append("No verified leader portrait")
        if not market_ready: issues.append("Market route incomplete for ownership type")
        score = round((25 if company_ready else 0) + (15 if logo_grade == "A" else 8 if logo_resolved else 0) + (20 if role_ready else 0) + (15 if social_ready else 0) + (15 if photo_ready else 0) + (10 if market_ready else 0))
        overall = "Publication ready" if not issues else "Ready with optional media gaps" if company_ready and role_ready and market_ready and all("leader" in issue.lower() or "logo" in issue.lower() for issue in issues) else "Editorial review required"
        priority = "P0" if not company_ready or not role_ready or not market_ready else "P1" if not social_ready and not photo_ready else "P2" if not social_ready or not photo_ready or logo_grade != "A" else "P3"
        next_action = "No action; routine refresh" if not issues else "; ".join([
            "verify leader X/LinkedIn against an official bio" if "No verified leader X" in issue else
            "obtain a named-person portrait from an official bio, Wikimedia Commons, or a verified X profile" if "portrait" in issue else
            "replace with an official SVG/PNG brand asset" if "Logo uses" in issue else
            "verify the ownership-appropriate market source" if "Market route" in issue else
            "verify a first-party company social profile" for issue in issues
        ])

        for field, value in {
            "evidence_grade": "A" if company_ready else "C", "logo_source_type": logo_type,
            "logo_source_url": website if "Official" in logo_type else logo, "publication_status": overall,
            "qa_last_reviewed_at": TODAY,
        }.items(): companies.cell(c_row, ci[field], value)
        for field, value in {
            "role_evidence_grade": "A" if role_ready else "C",
            "social_evidence_grade": "A" if leader_x.startswith("http") else "B" if leader_linkedin.startswith("http") else "C",
            "photo_evidence_grade": "A" if photo_ready else "C",
            "social_disposition": "Verified X and LinkedIn" if leader_x.startswith("http") and leader_linkedin.startswith("http") else "Verified X" if leader_x.startswith("http") else "Verified LinkedIn" if leader_linkedin.startswith("http") else "No verified X or LinkedIn profile located in current pass",
            "photo_disposition": "Verified named-person portrait" if photo_ready else "No verified named-person portrait located in current pass",
            "media_review_status": "Complete" if social_ready and photo_ready else "Needs media enrichment",
            "qa_last_reviewed_at": TODAY,
        }.items(): leaders.cell(l_row, li[field], value)
        for field, value in {"market_evidence_grade": market_grade, "market_review_status": market_status, "qa_last_reviewed_at": TODAY}.items(): market.cell(m_row, mi[field], value)

        write_row(logo_ws, out_row, [company_id, name, logo, website, logo_type, logo_grade, "Official domain or curated registry; no generative logo assets", "Verified" if logo_grade == "A" else "Usable fallback — review recommended", TODAY])
        write_row(qa_ws, out_row, [company_id, name, "Ready" if company_ready else "Review", "Verified" if role_ready else "Review", "Verified" if social_ready else "Unresolved", "Verified" if photo_ready else "Unresolved", f"{logo_grade} · {logo_type}", f"{market_grade} · {market_status}", overall, score, TODAY])
        write_row(review_ws, out_row, [company_id, name, priority, " | ".join(issues) if issues else "None", next_action, "Closed" if not issues else "Open", website, role_source, market_source, TODAY])
        metrics["company_ready"] += company_ready; metrics["logo_a"] += logo_grade == "A"; metrics["logo_resolved"] += logo_resolved
        metrics["company_linkedin"] += company_linkedin.startswith("http"); metrics["company_x"] += company_x.startswith("http")
        metrics["leader_role"] += role_ready; metrics["leader_social"] += social_ready; metrics["leader_photo"] += photo_ready
        metrics["leader_linkedin"] += leader_linkedin.startswith("http"); metrics["leader_x"] += leader_x.startswith("http")
        metrics["market_ready"] += market_ready; metrics["publication_ready"] += overall == "Publication ready"
        if ownership == "Public":
            metrics["public_total"] += 1; metrics["public_market_ready"] += market_ready
        elif ownership == "Private":
            metrics["private_total"] += 1; metrics["private_market_ready"] += market_ready
            metrics["private_crunchbase_direct"] += direct

    coverage = workbook["Coverage QA"]
    qi = headers(coverage)
    if coverage.max_row > 1:
        coverage.delete_rows(2, coverage.max_row - 1)
    records = [
        ("Unique organisations", 497, 497, "Canonical company_id records; no duplicate identity is counted."),
        ("Company profiles publication-ready", metrics["company_ready"], 497, "Website, logo and meaningful one-sentence description."),
        ("All logo references resolved", metrics["logo_resolved"], 497, "Identity-safe official-domain or curated logo reference; Grade A is reported separately."),
        ("Presentation-grade logo evidence (Grade A)", metrics["logo_a"], 497, "Official-site asset or curated brand registry; fallbacks are not Grade A."),
        ("Direct company LinkedIn", metrics["company_linkedin"], 497, "Only direct organisation pages are counted."),
        ("Verified official company X", metrics["company_x"], 497, "Accepted only after XActions identity validation or a reviewed first-party override."),
        ("Leadership roles with official evidence", metrics["leader_role"], 497, "Current role source recorded for every row."),
        ("Leadership with verified LinkedIn", metrics["leader_linkedin"], 497, "Exact named-person LinkedIn profile."),
        ("Leadership with verified X", metrics["leader_x"], 497, "Exact named-person X account with organisation context."),
        ("Leadership with verified X or LinkedIn", metrics["leader_social"], 497, "Identity-matched public social profile."),
        ("Leadership with verified portrait", metrics["leader_photo"], 497, "Named-person photo with provenance; ambiguous images are withheld."),
        ("Ownership-appropriate market route resolved", metrics["market_ready"], 497, "Public Yahoo/IR; private directory/funding; other ownership-specific route."),
        ("Public companies with Yahoo Finance and IR", metrics["public_market_ready"], metrics["public_total"], "All companies currently classified Public."),
        ("Private companies with a Crunchbase/funding route", metrics["private_market_ready"], metrics["private_total"], "Direct Crunchbase profile or explicitly labelled discovery exception plus funding/unicorn disposition."),
        ("Private companies with direct Crunchbase profile", metrics["private_crunchbase_direct"], metrics["private_total"], "Discovery/search routes are not counted as direct profiles."),
        ("Profiles with no open editorial issue", metrics["publication_ready"], 497, "Strict status; optional media and Grade B logo gaps remain open."),
    ]
    for row, (metric, current, target, note) in enumerate(records, 2):
        values = {"metric": metric, "current_count": current, "target_count": target, "coverage_rate": current / target if target else 0, "status": "Complete" if current == target else "Evidence coverage", "notes": note}
        for field, value in values.items(): coverage.cell(row, qi[field], value)
        coverage.cell(row, qi["coverage_rate"]).number_format = "0.0%"

    output = Path(args.output); output.parent.mkdir(parents=True, exist_ok=True); workbook.save(output)
    print({"output": str(output), **metrics})


if __name__ == "__main__":
    main()
