#!/usr/bin/env python3
"""Build the compact, source-aware company profiles used by the public map."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlparse

from openpyxl import load_workbook


EMPTY_VALUES = {
    "",
    "not applicable",
    "not publicly disclosed",
    "not found in reviewed sources",
    "no verified public profile located",
    "no verified public headshot located",
}


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return None if value.lower() in EMPTY_VALUES else value
    return value


def clean_url(value: Any) -> str | None:
    """Publish only explicit web addresses in fields rendered as links or media."""
    value = clean(value)
    if not isinstance(value, str):
        return None
    return value if value.lower().startswith(("http://", "https://")) else None


def platform_url(value: Any, hosts: tuple[str, ...]) -> str | None:
    value = clean_url(value)
    if not value:
        return None
    host = (urlparse(value).hostname or "").lower().removeprefix("www.")
    return value if host in hosts else None


def official_site_icon(website: Any) -> str | None:
    website = clean_url(website)
    if not website:
        return None
    host = urlparse(website).hostname
    return f"https://www.google.com/s2/favicons?domain={quote(host or '')}&sz=128" if host else None


def renderable_media_url(value: Any) -> str | None:
    """Convert Wikimedia file pages to image redirects usable by an img element."""
    value = clean_url(value)
    if not value:
        return None
    parsed = urlparse(value)
    host = (parsed.hostname or "").lower()
    fragment = unquote(parsed.fragment or "")
    path = unquote(parsed.path or "")
    candidate = fragment.removeprefix("/media/") if fragment.startswith("/media/") else path.rsplit("/", 1)[-1]
    for prefix in ("File:", "Berkas:"):
        if candidate.startswith(prefix):
            filename = candidate[len(prefix) :]
            special = "Istimewa:Redirect/file" if host == "id.wikipedia.org" else "Special:Redirect/file"
            return f"https://{host}/wiki/{special}/{quote(filename)}"
    return value


def curated_logo_url(value: Any) -> str | None:
    """Accept only independently curated brand assets with predictable identity."""
    value = renderable_media_url(value)
    if not value:
        return None
    parsed = urlparse(value)
    host = (parsed.hostname or "").lower().removeprefix("www.")
    if host in {"cdn.simpleicons.org", "api.iconify.design", "upload.wikimedia.org"}:
        return value
    return None


def deep_merge(target: dict[str, Any], patch: dict[str, Any]) -> None:
    for key, value in patch.items():
        if isinstance(value, dict) and isinstance(target.get(key), dict):
            deep_merge(target[key], value)
        else:
            target[key] = value


def rows_by_key(workbook, sheet_name: str, key: str) -> dict[str, dict[str, Any]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    output: dict[str, dict[str, Any]] = {}
    for cells in sheet.iter_rows(min_row=2):
        record = {header: clean(cell.value) for header, cell in zip(headers, cells)}
        record_key = clean(record.get(key))
        if record_key:
            output[str(record_key)] = record
    return output


def split_values(value: Any) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in str(value).split("|") if item.strip() and item.strip().lower() not in EMPTY_VALUES]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument(
        "--overrides",
        default=str(Path(__file__).resolve().parents[1] / "data" / "manual-profile-overrides.json"),
    )
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    companies = rows_by_key(workbook, "Companies DB", "company_id")
    leaders = rows_by_key(workbook, "Leadership DB", "company_id")
    markets = rows_by_key(workbook, "Market Data", "company_id")
    sources = rows_by_key(workbook, "Sources DB", "source_id")
    overrides_path = Path(args.overrides)
    overrides = json.loads(overrides_path.read_text(encoding="utf-8")) if overrides_path.exists() else {}

    profiles: list[dict[str, Any]] = []
    for company_id, company in companies.items():
        leader = leaders.get(company_id, {})
        market = markets.get(company_id, {})
        company_source = sources.get(str(company.get("primary_source_id") or ""), {})
        leader_source = sources.get(str(leader.get("source_id") or ""), {})

        longitude = clean(company.get("longitude"))
        latitude = clean(company.get("latitude"))
        location_precision = clean(company.get("location_precision")) or "Country-level placement"
        coordinates_verified = location_precision.lower().startswith("verified city")

        # Final publication corrections: never expose a mismatched person image or
        # force a startup-database identity onto a state-owned utility.
        if company_id == "boston-dynamics":
            leader["photo_url"] = None
            leader["photo_verification_status"] = "Removed after identity QA; official headshot required"
            company["hq_city"] = "Waltham, Massachusetts"
            company["data_notes"] = "Headquarters city and current CEO are source-verified; the previously mismatched leadership image is withheld pending an official headshot."
            longitude, latitude = -71.2356, 42.3765
            coordinates_verified = True
            location_precision = "Verified city"
        if company_id == "pln":
            market["crunchbase_url"] = None
            market["crunchbase_profile_status"] = "Not applicable to a state-owned enterprise"
            market["investor_relations_url"] = "https://web.pln.co.id/stakeholder/informasi-penunjang-pasar-modal"
            company["hq_city"] = "Jakarta"
            company["description_short"] = "PLN operates Indonesia's electricity generation, transmission and distribution infrastructure supporting the country's digital and AI economy."
            company["data_notes"] = "PLN is a state-owned enterprise headquartered in Jakarta; official corporate and capital-market sources replace the incorrect Crunchbase identity."
            longitude, latitude = 106.8456, -6.2088
            coordinates_verified = True
            location_precision = "Verified city"
        if company_id == "sandisk":
            company["hq_city"] = "Milpitas, California"
            company["data_notes"] = (
                "Sandisk's principal executive offices are at 951 Sandisk Drive, "
                "Milpitas, California; the map uses a city-level coordinate rather than a US centroid."
            )
            longitude, latitude = -121.8996, 37.4323
            coordinates_verified = True
            location_precision = "Verified city"
        if company_id == "eklipse":
            company["hq_city"] = "Singapore"
            company["hq_country"] = "Singapore"
            company["country_iso2"] = "SG"
            company["region"] = "Southeast Asia"
            company["data_notes"] = (
                "Eklipse is operated by Singapore-registered Main Spring Technology Pte Ltd; "
                "Indonesia and Vietnam are operating or development footprints rather than headquarters."
            )
            longitude, latitude = 103.8198, 1.3521
            coordinates_verified = True
            location_precision = "Verified city"
        if company_id == "pixel-ml":
            leader["linkedin_url"] = "https://www.linkedin.com/in/seanphan"
            leader["social_url"] = None
            leader["official_profile_url"] = "https://pixelml.com/about"
            leader["social_verification_status"] = "Verified via Pixel ML official About page"
            company["data_notes"] = (
                "Sean Phan's leadership identity and LinkedIn destination are verified from Pixel ML's official About page."
            )

        ownership = clean(company.get("ownership_status")) or "Ownership not classified"
        is_public = ownership == "Public"
        is_private = ownership == "Private"
        is_state_owned = ownership == "Government-owned"

        profile = {
            "id": company_id,
            "name": company.get("display_name"),
            "description": company.get("description_short"),
            "entityType": company.get("entity_type"),
            "ownershipStatus": ownership,
            "lifecycleStatus": company.get("lifecycle_status"),
            "foundedYear": company.get("founded_year"),
            "website": clean_url(company.get("official_website")),
            "linkedin": clean_url(company.get("company_linkedin")),
            "x": platform_url(company.get("company_x_url"), ("x.com", "twitter.com", "mobile.twitter.com")),
            "xAvatar": clean_url(company.get("company_x_avatar_url")),
            "logo": curated_logo_url(company.get("logo_url")) or official_site_icon(company.get("official_website")),
            "headquarters": {
                "city": company.get("hq_city"),
                "country": company.get("hq_country"),
                "countryIso2": company.get("country_iso2"),
                "region": company.get("region"),
                "longitude": longitude,
                "latitude": latitude,
                "precision": location_precision,
                "coordinatesVerified": coordinates_verified,
            },
            "aiPosition": {
                "primaryLayer": company.get("primary_layer"),
                "secondaryLayers": split_values(company.get("secondary_layers")),
                "primaryCapability": company.get("primary_capability"),
                "primaryVertical": company.get("primary_vertical"),
                "secondaryVerticals": split_values(company.get("secondary_verticals")),
                "aiIntensity": company.get("ai_intensity"),
                "whyIncluded": company.get("why_included"),
            },
            "leader": {
                "name": leader.get("full_name") or company.get("current_leader_name"),
                "position": leader.get("position") or company.get("current_leader_position"),
                "linkedin": platform_url(leader.get("linkedin_url"), ("linkedin.com",)) or platform_url(leader.get("social_url"), ("linkedin.com",)),
                "x": platform_url(leader.get("x_url"), ("x.com", "twitter.com", "mobile.twitter.com")),
                "xAvatar": clean_url(leader.get("x_avatar_url")),
                "photo": renderable_media_url(leader.get("photo_url")) or clean_url(leader.get("x_avatar_url")),
                "officialProfile": clean_url(leader.get("official_profile_url")),
                "roleSource": clean_url(leader_source.get("source_url")) or clean_url(leader.get("official_profile_url")),
                "photoSource": clean_url(leader.get("photo_source_url")) or (clean_url(leader.get("x_source_url")) if not clean_url(leader.get("photo_url")) else None),
                "socialStatus": leader.get("social_verification_status"),
                "xStatus": leader.get("x_verification_status"),
                "photoStatus": leader.get("photo_verification_status"),
            },
            "market": {
                "type": "public" if is_public else "private" if is_private else "state-owned" if is_state_owned else "other",
                "ticker": market.get("ticker"),
                "exchange": market.get("exchange"),
                "yahooFinance": clean_url(market.get("yahoo_finance_url")),
                "investorRelations": clean_url(market.get("investor_relations_url")),
                "crunchbase": clean_url(market.get("crunchbase_url")),
                "dealroom": clean_url(market.get("dealroom_url")),
                "latestRoundDate": market.get("latest_round_date"),
                "latestRoundType": market.get("latest_round_type"),
                "latestRoundAmountUsd": market.get("latest_round_amount_usd"),
                "totalFundingUsd": market.get("total_disclosed_funding_usd"),
                "latestValuationUsd": market.get("latest_disclosed_valuation_usd"),
                "valuationDate": market.get("valuation_date"),
                "unicornStatus": market.get("unicorn_status"),
                "fundingStatus": market.get("funding_history_status"),
            },
            "verification": {
                "status": company.get("verification_status"),
                "confidence": company.get("profile_confidence"),
                "lastVerifiedAt": company.get("last_verified_at"),
                "notes": company.get("data_notes"),
                "companySource": clean_url(company_source.get("source_url")) or clean_url(company.get("official_website")),
                "leaderSource": clean_url(leader_source.get("source_url")) or clean_url(leader.get("official_profile_url")),
            },
        }
        deep_merge(profile, overrides.get(company_id, {}))
        profiles.append(profile)

    profiles.sort(key=lambda item: (item.get("name") or "").lower())
    payload = {
        "meta": {
            "generatedAt": date.today().isoformat(),
            "count": len(profiles),
            "methodology": "Only source-aware identity, leadership, location and market fields are published. Unverified media and inapplicable ownership links are omitted.",
        },
        "companies": profiles,
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print({"output": str(output), "profiles": len(profiles)})


if __name__ == "__main__":
    main()
